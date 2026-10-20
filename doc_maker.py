"""
Заполнение шаблона Word и создание PDF.
Не зависит от Telegram: можно проверить отдельно — python doc_maker.py
"""
import re
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

PLACEHOLDER = re.compile(r'\{\{\s*(\w+)\s*\}\}')

# Шрифты с русскими буквами: сначала Windows, потом Linux
FONT_CANDIDATES = [
    (r'C:\Windows\Fonts\times.ttf', r'C:\Windows\Fonts\timesbd.ttf'),
    (r'C:\Windows\Fonts\arial.ttf', r'C:\Windows\Fonts\arialbd.ttf'),
    ('/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf',
     '/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf'),
    ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
     '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
]
_font_ready = False


def _register_font():
    global _font_ready
    if _font_ready:
        return
    for regular, bold in FONT_CANDIDATES:
        if Path(regular).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont('Main', regular))
            pdfmetrics.registerFont(TTFont('Main-Bold', bold))
            _font_ready = True
            return
    raise RuntimeError('Не найден шрифт с русскими буквами для PDF')


def _all_paragraphs(doc):
    """Абзацы документа, включая абзацы внутри таблиц."""
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def fill_docx(template_path, values, out_path):
    """Подставляет значения вместо {{полей}} и сохраняет новый .docx.
    Возвращает список полей, для которых не нашлось значения."""
    doc = Document(template_path)
    missing = set()

    def replace(match):
        key = match.group(1)
        if key in values:
            return str(values[key])
        missing.add(key)
        return match.group(0)

    for p in _all_paragraphs(doc):
        if '{{' not in p.text:
            continue
        new_text = PLACEHOLDER.sub(replace, p.text)
        # Word может разбить {{поле}} на куски, поэтому собираем текст
        # в первый кусок (с его оформлением), остальные очищаем
        runs = p.runs
        runs[0].text = new_text
        for r in runs[1:]:
            r.text = ''

    doc.save(out_path)
    return sorted(missing)


def docx_to_pdf(docx_path, pdf_path):
    """Делает PDF из заполненного .docx (текст, жирный, выравнивание)."""
    _register_font()
    doc = Document(docx_path)
    align_map = {WD_ALIGN_PARAGRAPH.CENTER: TA_CENTER,
                 WD_ALIGN_PARAGRAPH.RIGHT: TA_RIGHT,
                 WD_ALIGN_PARAGRAPH.JUSTIFY: TA_JUSTIFY}
    story = []
    for p in doc.paragraphs:
        text = p.text
        if not text.strip():
            story.append(Spacer(1, 5 * mm))
            continue
        bold = bool(p.runs) and all(r.bold for r in p.runs if r.text.strip())
        size = 14 if bold and p.alignment == WD_ALIGN_PARAGRAPH.CENTER else 12
        style = ParagraphStyle(
            'p', fontName='Main-Bold' if bold else 'Main', fontSize=size,
            leading=size * 1.4, spaceAfter=3 * mm,
            alignment=align_map.get(p.alignment, TA_LEFT))
        # несколько пробелов подряд в PDF схлопываются — сохраняем их
        safe = escape(text).replace('  ', '&nbsp; ')
        story.append(Paragraph(safe, style))

    pdf = SimpleDocTemplate(str(pdf_path), pagesize=A4,
                            leftMargin=25 * mm, rightMargin=15 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm)
    pdf.build(story)


def make_documents(template_path, values, out_folder, base_name):
    """Готовит .docx и .pdf. Возвращает пути к двум файлам."""
    out_folder = Path(out_folder)
    out_folder.mkdir(parents=True, exist_ok=True)
    values = dict(values)
    values.setdefault('today', datetime.now().strftime('%d.%m.%Y'))
    values.setdefault('doc_no', datetime.now().strftime('%y%m%d-%H%M'))

    docx_path = out_folder / f'{base_name}.docx'
    pdf_path = out_folder / f'{base_name}.pdf'
    missing = fill_docx(template_path, values, docx_path)
    if missing:
        print('В шаблоне есть поля без значений:', ', '.join(missing))
    docx_to_pdf(docx_path, pdf_path)
    return docx_path, pdf_path


if __name__ == '__main__':
    # Быстрая проверка без Telegram
    test = {'city': 'Москва', 'client_name': 'Иванов Иван Иванович',
            'client_phone': '+7 900 123-45-67', 'service': 'настройка компьютера',
            'price': '3 000', 'deadline': '25.10.2026'}
    files = make_documents('templates/Договор_оказания_услуг.docx', test,
                           'output', 'Проверка_договора')
    print('Готово:', *files, sep='\n  ')
