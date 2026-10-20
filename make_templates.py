"""
Создаёт примеры шаблонов в папке templates.
Запустить один раз: python make_templates.py

Свои шаблоны можно сделать прямо в Word: пишете текст документа,
а вместо данных ставите поля в двойных фигурных скобках, например {{client_name}}.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

FOLDER = Path(__file__).parent / 'templates'


def new_doc():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    return doc


def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def para(doc, text, align=None):
    p = doc.add_paragraph(text)
    if align == 'right':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == 'justify':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def contract():
    doc = new_doc()
    title(doc, 'ДОГОВОР ОКАЗАНИЯ УСЛУГ № {{doc_no}}')
    para(doc, 'г. {{city}}                                                                                     {{today}}')
    para(doc, '{{client_name}}, именуемый(ая) в дальнейшем «Заказчик», с одной стороны, '
              'и Исполнитель, с другой стороны, заключили настоящий договор о нижеследующем:', 'justify')
    para(doc, '1. ПРЕДМЕТ ДОГОВОРА').runs[0].bold = True
    para(doc, '1.1. Исполнитель обязуется оказать Заказчику следующие услуги: {{service}}.', 'justify')
    para(doc, '1.2. Срок оказания услуг: до {{deadline}}.', 'justify')
    para(doc, '2. СТОИМОСТЬ И ПОРЯДОК РАСЧЁТОВ').runs[0].bold = True
    para(doc, '2.1. Стоимость услуг составляет {{price}} рублей.', 'justify')
    para(doc, '2.2. Оплата производится после подписания акта выполненных работ.', 'justify')
    para(doc, '3. КОНТАКТЫ СТОРОН').runs[0].bold = True
    para(doc, 'Заказчик: {{client_name}}, тел. {{client_phone}}')
    para(doc, '')
    para(doc, 'Заказчик ____________________          Исполнитель ____________________')
    doc.save(FOLDER / 'Договор_оказания_услуг.docx')


def vacation():
    doc = new_doc()
    para(doc, 'Руководителю {{company}}', 'right')
    para(doc, '{{director}}', 'right')
    para(doc, 'от {{position}}', 'right')
    para(doc, '{{employee}}', 'right')
    para(doc, '')
    title(doc, 'ЗАЯВЛЕНИЕ')
    para(doc, 'Прошу предоставить мне ежегодный оплачиваемый отпуск с {{date_from}} '
              'продолжительностью {{days}} календарных дней.', 'justify')
    para(doc, '')
    para(doc, '{{today}}                                              ____________ / {{employee}}')
    doc.save(FOLDER / 'Заявление_на_отпуск.docx')


def receipt():
    doc = new_doc()
    title(doc, 'РАСПИСКА В ПОЛУЧЕНИИ ДЕНЕЖНЫХ СРЕДСТВ')
    para(doc, 'г. {{city}}                                                                                     {{today}}')
    para(doc, 'Я, {{borrower}}, получил(а) от {{lender}} денежные средства в размере '
              '{{amount}} рублей.', 'justify')
    para(doc, 'Обязуюсь вернуть указанную сумму в срок до {{return_date}}.', 'justify')
    para(doc, 'Деньги получены полностью, претензий не имею.', 'justify')
    para(doc, '')
    para(doc, '____________ / {{borrower}}')
    doc.save(FOLDER / 'Расписка.docx')


if __name__ == '__main__':
    FOLDER.mkdir(exist_ok=True)
    contract()
    vacation()
    receipt()
    print('Шаблоны созданы в папке templates')
